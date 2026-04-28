"""Convert report/report.md to a formatted Word .docx file."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

REPORT_MD = Path(__file__).parent / "report.md"
OUT_DOCX = Path(__file__).parent / "NLP_Project_Report.docx"


def set_font(run, name="Times New Roman", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    sizes = {1: 16, 2: 13, 3: 12}
    set_font(run, name="Times New Roman", size=sizes.get(level, 12), bold=True)
    return p


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Split on inline italics/bold markers
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = p.add_run(tok[2:-2])
            set_font(run, size=11, bold=True)
        elif tok.startswith("*") and tok.endswith("*"):
            run = p.add_run(tok[1:-1])
            set_font(run, size=11, italic=True)
        elif tok.startswith("`") and tok.endswith("`"):
            run = p.add_run(tok[1:-1])
            set_font(run, name="Consolas", size=10)
        else:
            run = p.add_run(tok)
            set_font(run, size=11)


def add_monospace_block(doc, text):
    """Render a preformatted code/diagram block in Consolas."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(text)
    set_font(run, name="Consolas", size=9)


def add_results_table(doc):
    """Add Table 1 — retrieval metrics."""
    table = doc.add_table(rows=5, cols=4)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    headers = ["Method", "Recall@5", "MRR@10", "nDCG@10"]
    rows = [
        ["BM25", "0.6047", "0.4460", "0.5159"],
        ["Dense (MiniLM-L6)", "0.3635", "0.2360", "0.2888"],
        ["Hybrid (α = 0.6)", "0.6047", "0.4594", "0.5178"],
        ["Hybrid + Rerank", "0.6700", "0.5118", "0.5705"],
    ]

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(h)
        set_font(run, size=10.5, bold=True)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if c > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = para.add_run(val)
            bold = (r == 4)
            set_font(run, size=10.5, bold=bold)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_title_block(doc):
    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Domain-Specific Question Answering on Indian Policy Corpora "
        "using Hybrid Retrieval and a Custom Decoder-Only Transformer"
    )
    set_font(run, size=16, bold=True)

    # Author line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Krish Saini (230708)")
    set_font(run, size=12, bold=True)

    # Mentor + institution line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Mentor: Dr. Atul Mishra")
    set_font(run, size=11, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BML Munjal University — Natural Language Processing Project")
    set_font(run, size=11, italic=True)

    # Spacer
    doc.add_paragraph()


def render(md: str, doc: Document) -> None:
    add_title_block(doc)

    # Parse markdown into blocks
    lines = md.split("\n")
    i = 0
    in_code = False
    code_buf: list[str] = []
    skip_title_block = True  # skip front-matter until first ## heading

    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip()

        if stripped.startswith("```"):
            if in_code:
                add_monospace_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Skip the YAML-ish title block
        if skip_title_block:
            if stripped.startswith("## "):
                skip_title_block = False
            else:
                i += 1
                continue

        if stripped.startswith("## "):
            add_heading(doc, stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), level=2)
            i += 1
            continue

        # Table 1 — detect the retrieval table and render via python-docx
        if stripped.startswith("| Method"):
            # consume the markdown table (header + separator + 4 rows)
            add_results_table(doc)
            # Add centered caption
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            crun = cap.add_run("Table 1. Retrieval metrics on 597 domain QA pairs.")
            set_font(crun, size=10, italic=True)
            # skip rows in source
            while i < len(lines) and lines[i].startswith("|"):
                i += 1
            continue

        # Skip the bold caption line preceding the table (already rendered below)
        if stripped.startswith("**Table 1"):
            i += 1
            continue

        if stripped.startswith("---"):
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Collect paragraph text across wrapped lines
        buf = [stripped]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].startswith(("#", "|", "```", "---"))
        ):
            buf.append(lines[i].strip())
            i += 1
        add_body_paragraph(doc, " ".join(buf))


def main() -> None:
    md = REPORT_MD.read_text(encoding="utf-8")
    doc = Document()

    # Set default margins
    for section in doc.sections:
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    render(md, doc)
    doc.save(OUT_DOCX)
    print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    main()
